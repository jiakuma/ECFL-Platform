import json
from collections import Counter
from openpyxl import load_workbook

path = r"D:\desktop\解放军联邦学习\元数据-食管 (2).xlsx"
wb = load_workbook(path, data_only=True)
ws = wb.worksheets[0]

rows = list(ws.iter_rows(min_row=2, values_only=True))
fields = []
skipped = []
for r in rows:
    cn, en, abbr, definition, dtype, domain, constraint, maxocc, category = (list(r) + [None] * 9)[:9]
    cn = (cn or "").strip() if isinstance(cn, str) else ""
    if not cn:
        continue
    if cn.endswith("：") or cn.endswith(":"):
        skipped.append(cn)
        continue
    fields.append({
        "cn": cn,
        "en": (en or "").strip() if isinstance(en, str) else "",
        "abbr": (abbr or "").strip() if isinstance(abbr, str) else "",
        "dtype": (dtype or "").strip() if isinstance(dtype, str) else "",
        "domain": (domain or "").strip() if isinstance(domain, str) else "",
        "constraint": (constraint or "").strip() if isinstance(constraint, str) else "",
        "category": (category or "").strip() if isinstance(category, str) else "",
    })

print("section headers skipped:", skipped)
print("total standard fields:", len(fields))
c = Counter(f["category"] for f in fields)
for k, v in c.most_common():
    print(f"  {k}: {v}")
with open("scripts/meta_fields.json", "w", encoding="utf-8") as f:
    json.dump(fields, f, ensure_ascii=False, indent=1)
