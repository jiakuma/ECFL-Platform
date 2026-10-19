import sys
from openpyxl import load_workbook

path = r"D:\desktop\解放军联邦学习\元数据-食管 (2).xlsx"
wb = load_workbook(path, data_only=True)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print("\n=== SHEET:", ws.title, "dims:", ws.dimensions, "max_row:", ws.max_row, "max_col:", ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200), values_only=True):
        cells = ["" if c is None else str(c).replace("\n", " ").strip() for c in row]
        if any(cells):
            print(" | ".join(cells))
